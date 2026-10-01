# src/council/sources.py
"""
Background and documents for a council run.

  --context  inline text, or a path to a file whose text becomes the background
  --docs     files and/or folders attached to this one question

Text is extracted up front, in code, so the council always sees the material
(it doesn't depend on the moderator deciding to look). Short documents go into
the advisors' brief word for word; long ones are excerpted to stay within
COUNCIL_DOC_BUDGET characters, and the moderator can read them in full with
the read_document tool.

Readable: .md .txt .csv .tsv .json .yaml .xml .html (plain text), .pdf, .docx, .xlsx
Not readable: scanned PDFs (images, no text layer), legacy .doc/.xls, images.
"""
from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Type

from crewai.tools import BaseTool
from pydantic import BaseModel, Field, PrivateAttr

log = logging.getLogger("council")

TEXT_TYPES = {".md", ".markdown", ".txt", ".text", ".csv", ".tsv", ".json", ".jsonl",
              ".yaml", ".yml", ".xml", ".html", ".htm", ".rst", ".log", ".ini", ".toml"}
DOC_TYPES = {".pdf", ".docx", ".xlsx", ".xlsm"}
SUPPORTED = TEXT_TYPES | DOC_TYPES
LEGACY = {".doc": "save it as .docx", ".xls": "save it as .xlsx", ".ppt": "export it to PDF",
          ".pptx": "export it to PDF"}

MAX_SHEET_ROWS = 2000     # per Excel sheet
TOOL_PAGE = 20_000        # characters returned per read_document call


def workspace_dir() -> Path:
    """Shared folder the moderator browses when no documents are attached."""
    return Path(os.getenv("COUNCIL_WORKSPACE", "./workspace"))


WORKSPACE_HINT = "the moderator will check the shared workspace folder instead"


def doc_budget() -> int:
    """Characters of document text included verbatim in the brief (~4 characters per token)."""
    return int(os.getenv("COUNCIL_DOC_BUDGET", "30000"))


class SourceError(ValueError):
    """A file the user named can't be used. Raised before any model is called."""


# ------------------------------------------------------------------ extraction

def extract_text(path: Path) -> str:
    """Plain text of a supported file. Raises SourceError if it can't be read."""
    ext = path.suffix.lower()
    try:
        if ext in TEXT_TYPES:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
        elif ext == ".pdf":
            text = _pdf_text(path)
        elif ext == ".docx":
            text = _docx_text(path)
        elif ext in (".xlsx", ".xlsm"):
            text = _xlsx_text(path)
        elif ext in LEGACY:
            raise SourceError(f"{path.name}: {ext} files can't be read; {LEGACY[ext]}.")
        else:
            raise SourceError(f"{path.name}: unsupported file type '{ext or '(none)'}'.")
    except SourceError:
        raise
    except Exception as e:  # corrupt, password-protected, locked by Excel...
        raise SourceError(f"{path.name}: couldn't be read ({type(e).__name__}: {e}).") from e

    text = text.strip()
    if not text:
        hint = " It may be a scanned image; run it through OCR first." if ext == ".pdf" else ""
        raise SourceError(f"{path.name}: no readable text.{hint}")
    return text


def _pdf_text(path: Path) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(f"[page {i}]\n{text}")
    return "\n\n".join(parts)


def _docx_text(path: Path) -> str:
    import docx
    from docx.table import Table

    lines = []
    for block in docx.Document(str(path)).iter_inner_content():  # paragraphs and tables, in order
        if isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("")
        elif block.text.strip():
            style = (block.style.name if block.style is not None else "") or ""
            prefix = "#" * int(style[-1]) + " " if style.startswith("Heading") and style[-1:].isdigit() else ""
            lines.append(prefix + block.text.strip())
    return "\n".join(lines)


def _xlsx_text(path: Path) -> str:
    import openpyxl

    # data_only: formula cells give their last calculated value (blank if the
    # file was produced by a program that never calculated them).
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            rows, clipped = [], False
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    clipped = True
                    break
                cells = ["" if v is None else str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                note = f"\n[sheet clipped at {MAX_SHEET_ROWS} rows]" if clipped else ""
                parts.append(f"## Sheet: {ws.title}\n" + "\n".join(rows) + note)
    finally:
        wb.close()
    return "\n\n".join(parts)


# ------------------------------------------------------------------ documents

@dataclass
class Document:
    path: Path
    name: str          # how the council refers to it
    text: str

    @property
    def chars(self) -> int:
        return len(self.text)


@dataclass
class Skipped:
    path: Path
    reason: str


def _is_hidden(part: str) -> bool:
    return part.startswith((".", "~$"))  # dotfiles, Office lock files


def _display_names(paths: list[Path]) -> list[str]:
    """Short names, lengthened only where two files would otherwise share one."""
    cwd = Path.cwd().resolve()

    def rel(p: Path) -> str:
        try:
            return p.relative_to(cwd).as_posix()
        except ValueError:
            return p.as_posix()

    base = [p.name for p in paths]
    return [b if base.count(b) == 1 else rel(p) for b, p in zip(base, paths)]


def load_documents(specs: list[str] | None) -> tuple[list[Document], list[Skipped]]:
    """
    Resolve --docs arguments (files or folders) into extracted documents.

    A path that doesn't exist is an error (raised before any model call).
    Files that can't be read are skipped and reported, so one bad PDF doesn't
    cancel the run.
    """
    files: list[Path] = []
    skipped: list[Skipped] = []
    for spec in specs or []:
        p = Path(spec).expanduser()
        if not p.exists():
            hint = (" If that's your question, put the question before --docs."
                    if " " in spec.strip() else "")
            raise SourceError(f"Not found: '{spec}'.{hint}")
        if p.is_dir():
            for f in sorted(p.rglob("*")):
                if not f.is_file() or any(_is_hidden(x) for x in f.relative_to(p).parts):
                    continue
                if f.suffix.lower() in SUPPORTED:
                    files.append(f)
                else:
                    skipped.append(Skipped(f, f"unsupported type '{f.suffix or '(none)'}'"))
        else:
            files.append(p)

    seen, unique = set(), []
    for f in files:
        key = f.resolve()
        if key not in seen:
            seen.add(key)
            unique.append(f)

    docs = []
    for f, name in zip(unique, _display_names(unique)):
        try:
            docs.append(Document(path=f, name=name, text=extract_text(f)))
        except SourceError as e:
            skipped.append(Skipped(f, str(e).split(": ", 1)[-1]))
    return docs, skipped


def load_context(value: str | None) -> tuple[str, str | None]:
    """
    --context value -> (background text, source file name or None).

    An existing file path is read; anything else is taken as the text itself.
    A single word ending in a known extension (e.g. "notes.md") that doesn't
    exist is almost certainly a typo'd path, so it's an error rather than
    silently becoming the background.
    """
    value = (value or "").strip()
    if not value:
        return "", None
    looks_like_path = "\n" not in value and len(value) < 500
    if looks_like_path and Path(value).expanduser().is_file():
        p = Path(value).expanduser()
        text = extract_text(p)
        limit = doc_budget()
        if len(text) > limit:
            log.warning("Background file %s is long; using its first %s characters. "
                        "Attach long material with --docs instead.", p.name, f"{limit:,}")
            text = text[:limit] + "\n[... background truncated]"
        return text, p.name
    ext = Path(value).suffix.lower()
    if looks_like_path and " " not in value and (ext in SUPPORTED or ext in LEGACY):
        raise SourceError(f"Context file not found: '{value}'.")
    return value, None


# ------------------------------------------------------------------ prompt text

def _allocate(docs: list[Document], budget: int) -> dict[str, int]:
    """Characters per document: small ones in full, the rest share what's left evenly."""
    alloc, remaining, left = {}, budget, sorted(docs, key=lambda d: d.chars)
    while left:
        share = remaining // len(left)
        d = left.pop(0)
        alloc[d.name] = min(d.chars, share)
        remaining -= alloc[d.name]
    return alloc


def prompt_inputs(docs: list[Document] | None) -> dict[str, str]:
    """
    The document placeholders used in tasks.yaml:
      {documents}          the documents' text (whole, or excerpted to the budget)
      {documents_index}    one line per document
      {document_guidance}  what the moderator should do about documents
    """
    docs = docs or []
    if not docs:
        return {
            "documents": "(no documents attached to this question)",
            "documents_index": "(none)",
            "document_guidance": (
                "No documents were attached to this question. Check the shared workspace "
                "with the workspace tools, read any files relevant to this question with "
                "read_document, and name the ones you used (or say none were relevant)."
            ),
        }

    alloc = _allocate(docs, doc_budget())
    blocks, index, excerpted = [], [], False
    for d in docs:
        n = alloc[d.name]
        if n >= d.chars:
            body, status = d.text, "included in full"
        else:
            excerpted = True
            body = (d.text[:n] + f"\n[... excerpt: first {n:,} of {d.chars:,} characters. "
                    f"The full text was read by the moderator.]")
            status = f"excerpted to the first {n:,}"
        blocks.append(f'<document name="{d.name}">\n{body}\n</document>')
        index.append(f"- {d.name} ({d.chars:,} characters; {status})")

    guidance = "All attached documents are included in full above."
    if excerpted:
        guidance = ("Some documents above are excerpts. Read those in full with the "
                    "read_document tool before summarizing, so nothing important is missed.")
    return {
        "documents": "\n\n".join(blocks),
        "documents_index": "\n".join(index),
        "document_guidance": guidance,
    }


# ------------------------------------------------------------------ tool

class ReadDocumentInput(BaseModel):
    name: str = Field(description="The document's name or path, exactly as listed.")
    start: int = Field(default=0, description="Character offset to start from, for long documents.")


class ReadDocumentTool(BaseTool):
    """
    Reads PDF, Word, Excel and text files, scoped to either the documents
    attached to this run or the shared workspace folder. It can't read
    anything else on the computer.
    """
    name: str = "read_document"
    description: str = (
        "Reads the full text of a document (PDF, Word, Excel or text file). "
        f"Returns up to {TOOL_PAGE:,} characters per call; for longer documents, "
        "call again with the 'start' value it gives you."
    )
    args_schema: Type[BaseModel] = ReadDocumentInput
    _docs: dict[str, Document] = PrivateAttr(default_factory=dict)
    _root: Path | None = PrivateAttr(default=None)

    @classmethod
    def for_documents(cls, docs: list[Document]) -> "ReadDocumentTool":
        tool = cls(description=cls.model_fields["description"].default
                   + " Available documents: " + ", ".join(d.name for d in docs) + ".")
        for d in docs:
            for key in {d.name, d.path.name, str(d.path), d.path.as_posix()}:
                tool._docs[key.lower()] = d
        return tool

    @classmethod
    def for_folder(cls, root: str | Path) -> "ReadDocumentTool":
        tool = cls()
        tool._root = Path(root).resolve()
        return tool

    def _find(self, name: str) -> tuple[str, str]:
        key = name.strip().strip("'\"")
        if self._docs:
            d = self._docs.get(key.lower()) or self._docs.get(Path(key).name.lower())
            if d is None:
                raise SourceError(f"No attached document called '{name}'.")
            return d.name, d.text
        if self._root is None:
            raise SourceError("No documents are available.")
        p = Path(key)
        p = (p if p.is_absolute() else self._root / p).resolve()
        if self._root not in p.parents and p != self._root:
            raise SourceError(f"'{name}' is outside the workspace folder.")
        if not p.is_file():
            raise SourceError(f"No file called '{name}' in the workspace.")
        return p.relative_to(self._root).as_posix(), extract_text(p)

    def _run(self, name: str, start: int = 0) -> str:
        try:
            label, text = self._find(name)
        except SourceError as e:
            return f"Error: {e}"
        start = max(0, int(start))
        end = min(len(text), start + TOOL_PAGE)
        more = (f"\n\n[characters {start:,}-{end:,} of {len(text):,}; "
                f"call again with start={end} for more]") if end < len(text) else ""
        return f"{label}:\n\n{text[start:end]}{more}"
